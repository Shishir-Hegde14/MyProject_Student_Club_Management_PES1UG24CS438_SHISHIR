from pathlib import Path
import csv
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
IDENTITY='Shishir Hegde | SRN PES1UG24CS438 | Section 5H'
PROJECT='Student Club Event Ticketing and Budget Portal'


def doc(title,landscape=False):
    d=Document()
    d.core_properties.author='Shishir Hegde'
    d.core_properties.title=title
    s=d.sections[0]
    s.page_width=Inches(11 if landscape else 8.5)
    s.page_height=Inches(8.5 if landscape else 11)
    if landscape: s.orientation=WD_ORIENT.LANDSCAPE
    s.top_margin=s.bottom_margin=Inches(.65)
    s.left_margin=s.right_margin=Inches(.65)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=d.styles[name]
        style.font.name='Arial'
        style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt({'Normal':11,'Title':17,'Heading 1':14,'Heading 2':12,'Heading 3':11}[name])
        style.paragraph_format.space_after=Pt(7)
    for style in d.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    d.styles['Normal'].paragraph_format.line_spacing=1.08
    d.add_paragraph(title,'Title')
    d.add_paragraph(PROJECT)
    d.add_paragraph(IDENTITY)
    foot=s.footer.paragraphs[0]
    foot.text='PES1UG24CS438 | 5H                                      '
    foot.runs[0].font.size=Pt(9)
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); foot._p.append(f)
    return d


def p(d,t): return d.add_paragraph(t)
def h(d,t): return d.add_paragraph(t,'Heading 1')
def page(d): d.add_page_break()


def table(d,headers,rows,widths,font=10.5):
    t=d.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for cell,w in zip(t.rows[0].cells,widths): cell.width=Inches(w)
    for cell,label in zip(t.rows[0].cells,headers): cell.text=label
    for row in rows:
        cells=t.add_row().cells
        for cell,value,w in zip(cells,row,widths): cell.text=str(value); cell.width=Inches(w)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        keep=OxmlElement('w:cantSplit');trpr.append(keep)
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders)
            margin=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margin.append(e)
            props.append(margin)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.line_spacing=1
                for run in para.runs:
                    run.font.name='Arial';run.font.size=Pt(font);run.font.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
    p(d,'')
    return t


def save(d,path):
    for border in list(d.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    out=ROOT/path
    d.save(out)
    return out


REQUIREMENTS=[
('FR-001','Functional','High','The system shall route submitted budget requests through Faculty Coordinator, Finance Officer and Dean in that order. Each decision shall record the authenticated reviewer, stage, decision and time. Rejection shall stop the workflow and require a reason.',
 'Pass: all three approvals occur in order and appear in the audit record. An early Dean approval, repeated sign-off or approval after rejection is refused. Fail: any stage is skipped or an unauthorised decision changes the status.',
 'This gives each office a chance to check a request before funds are committed.'),
('FR-002','Functional','High','The system shall let a Club Lead create and edit their own draft proposal with an event name, date, venue, capacity and itemised budget. Submission shall validate required fields and the club allocation, then reserve the requested amount.',
 'Pass: complete, affordable proposals enter Faculty pending. Missing fields, past dates, invalid amounts and requests above the remaining allocation stay out of the queue. The lead can reduce a draft budget and try again. Fail: invalid data or overspending is accepted.',
 'Reviewers need complete event details. Reserving submitted budgets prevents two pending requests from using the same money.'),
('FR-003','Functional','High','The system shall issue a unique, signed QR ticket when a student registers for an approved, unexpired event with space available. The ticket shall identify the event and attendee. Repeat registration shall return the existing ticket.',
 'Pass: different attendees receive different tickets; the same attendee receives one ticket per event. Draft, rejected, expired and full events reject registration. Fail: duplicate tickets or registrations above capacity are created.',
 'Each registration needs a verifiable ticket, and the number of attendees must stay within the venue capacity.'),
('FR-004','Functional','High','The system shall allow check-in staff to validate ticket text read from a QR code, accept an unused valid ticket once and mark it used. It shall reject altered, malformed, expired, wrong-event and previously used tickets.',
 'Pass: the first scan is accepted and a second is rejected. Concurrent scans of the same ticket produce exactly one acceptance. Invalid input returns a rejection without a server error. Fail: an invalid ticket or repeat entry is accepted.',
 'One attendee should not be able to share a ticket to admit several people.'),
('FR-005','Functional','Medium','The system shall show a Club Lead the current approval status, registration count and remaining allocation for their own events. The dashboard shall refresh status and counts within 5 seconds of a committed change during normal connected use.',
 'Pass: an approval or registration is reflected on the open dashboard within 5 seconds; another lead cannot view the event. Fail: the display remains stale or exposes another lead\'s proposal.',
 'The lead should be able to follow an event without contacting each approving office.'),
('NFR-001','Performance and security','High','The check-in endpoint shall validate entry authenticity in under 100 ms under the agreed peak load. It shall require staff authorisation, verify the ticket signature and prevent duplicate entry.',
 'Measure the complete HTTP validation request on the target server; use 20 simultaneous check-in clients as the initial peak-load assumption. Every request must stay below 100 ms, and tampering, role and replay tests must pass. Local results are reported separately from a production acceptance decision.',
 'Slow validation creates queues, while weak validation allows forged or reused tickets.'),
('NFR-002','Scalability and usability','Medium','The portal shall support at least 500 concurrent users during registration periods, with page loads under 2 seconds, and remain usable on mobile browsers.',
 'Test 500 concurrent authenticated users and measure browser page-load time below 2 seconds. At a 390-pixel viewport, forms must remain readable and usable; wide tables may scroll within their container. A server-only load test does not by itself prove browser-load performance.',
 'Registration traffic can rise quickly after an announcement, and students often use phones.')]


def requirements():
    d=doc('Requirements and Use Case Specification',True)
    p(d,'This is the revised requirements record for the individual project. It keeps the five functional and two non-functional requirement IDs from Lab 1. The original submission is retained without edits in the originals folder.')
    h(d,'Functional requirements')
    table(d,['ID','Type and priority','Description','Acceptance criteria','Rationale'],[(a,b+'\n'+c,e,f,g) for a,b,c,e,f,g in REQUIREMENTS[:3]],[.85,1.15,2.9,3.1,1.7])
    page(d);h(d,'Functional requirements continued')
    table(d,['ID','Type and priority','Description','Acceptance criteria','Rationale'],[(a,b+'\n'+c,e,f,g) for a,b,c,e,f,g in REQUIREMENTS[3:5]],[.85,1.15,2.9,3.1,1.7])
    h(d,'Clarifications used in the prototype')
    p(d,'A digital sign-off means a recorded decision by an authenticated reviewer, not a certificate-based digital signature. Money is stored as integer paise. Pending and approved requests reserve allocation; rejection releases it. The prototype has one configured allocation period and does not automatically renew annual budgets.')
    p(d,'Students register themselves. Campus Admin is separated into three reviewer roles and check-in staff so that one ordinary account cannot skip approval stages. These roles make the original requirements explicit; they do not add payment collection or university account integration.')
    page(d);h(d,'Non-functional requirements')
    table(d,['ID','Type and priority','Description','Acceptance criteria','Rationale'],[(a,b+'\n'+c,e,f,g) for a,b,c,e,f,g in REQUIREMENTS[5:]],[.85,1.15,2.9,3.1,1.7])
    p(d,'The performance values are acceptance targets, not claims about an unmeasured deployment. The local test method and observed results are in 6-Testing. Hardware scanner tests, real mobile-device tests and production load acceptance are still separate checks.')
    page(d);h(d,'Use case UC01 Request budget allocation')
    p(d,'Primary actor: Club Lead. Supporting actor: Faculty Coordinator, who receives the request. Related requirements: FR-001, FR-002 and FR-005.')
    p(d,'Preconditions: the Club Lead is logged in, belongs to a club and has saved a complete draft proposal for today or a future date. The club has a configured allocation.')
    p(d,'Postconditions on success: the request is Faculty pending, its amount is reserved against the club allocation, and the lead can see the new status. On failure: the proposal stays a draft and no money is reserved.')
    h(d,'Main success scenario')
    for s in ['1. The Club Lead opens a draft proposal and checks the event details and budget items.','2. The Club Lead selects Request budget allocation.','3. The system confirms ownership, draft status and the event date.','4. In one transaction, the system checks that the request fits within the remaining allocation.','5. The system changes the status to Faculty pending and reserves the amount.','6. The system displays confirmation. The Faculty Coordinator can now review the request.']:
        p(d,s)
    h(d,'Alternate flow A1 Budget exceeds the allocation')
    p(d,'At step 4, the request exceeds the remaining allocation. The system shows the remaining amount and leaves the proposal as a draft. The lead selects Edit draft, reduces the budget and saves it, then returns to step 2. If the lead leaves the page, the unchanged draft remains available.')
    save(d,Path('1-Requirements-Engineering/requirements-and-use-case.docx'))
    with (ROOT/'1-Requirements-Engineering/requirements.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.writer(f);w.writerow(['ID','Type','Priority','Description','Acceptance criteria','Rationale']);w.writerows(REQUIREMENTS)


def rtm():
    d=doc('Requirements Traceability Matrix',True)
    p(d,'Each requirement is linked to its use case, logical component, implementation and verification. Test names are defined in 6-Testing/test_portal.py. A test link means the behaviour was exercised; it does not turn a local benchmark into production acceptance.')
    rows=[
    ('FR-001','UC02 Review budget','Event and Budget','Portal.submit; Portal.review; approvals table','TC03, TC04, TC11','Passed'),
    ('FR-002','UC01 Request budget; create or edit proposal','Event and Budget','Portal.create_event; Portal.submit; budget_items','TC01, TC02, TC05, TC13, TC19','Passed'),
    ('FR-003','UC03 Register for event; generate QR ticket','Ticket Service','Portal.register; Portal.token; ticket_svg','TC06, TC07, TC08, TC16, TC18','Passed'),
    ('FR-004','UC04 Scan and validate ticket','Check-in Service','Portal.checkin; conditional ticket update','TC09, TC10, TC11, TC12, TC17','Passed; hardware scan not tested'),
    ('FR-005','UC05 Track event status','Dashboard Service','Portal.dashboard; /api/status; dashboard.js','TC11, TC14; browser walkthrough','API passed; polling interval 3 seconds'),
    ('NFR-001','UC04 Scan and validate ticket','Authentication; Check-in; SQLite','Role checks; HMAC; atomic update','TC09-TC12, TC15, TC17; benchmark.py','Security cases passed; see measured performance limits'),
    ('NFR-002','All browser use cases','Web UI; application; SQLite','Responsive CSS; status polling; HTTP routes','500-user benchmark; mobile viewport review','Partial verification; real-device and browser-load acceptance pending')]
    table(d,['Requirement','Use case','Component','Implementation','Verification','Status'],rows,[.75,1.45,1.3,2.2,2.1,1.9])
    save(d,Path('1-Requirements-Engineering/traceability-matrix.docx'))
    with (ROOT/'1-Requirements-Engineering/traceability-matrix.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.writer(f);w.writerow(['Requirement','Use case','Component','Implementation','Verification','Status']);w.writerows(rows)


def architecture():
    d=doc('Architecture Justification')
    h(d,'Architecture selection')
    p(d,'I chose a layered architecture for the Student Club Event Ticketing and Budget Portal. The browser handles the forms and display, the application handles the event rules, and a database stores the records. The prototype runs as one Flask application with SQLite. It is also client-server in deployment: browsers send requests to one application server.')
    h(d,'Reasons for the choice')
    p(d,'First, the portal has clear groups of work: proposals, approvals, tickets and check-in. Keeping these rules separate from the pages makes it easier to change a form without changing the approval process. It also makes the important rules easier to test.')
    p(d,'Second, this is a small individual project. A single application is easier to run and debug than several services. Approval records and budget reservations can share one database transaction, so a half-completed request is less likely to leave inconsistent data.')
    h(d,'Security advantage')
    p(d,'The browser cannot directly update the database. The application checks the logged-in role before changing a request or validating a ticket. Signed QR data makes changes to the ticket detectable, and a conditional database update prevents a second entry with the same ticket.')
    h(d,'Performance benefit and limitation')
    p(d,'The business components call each other locally, avoiding extra network calls between services. Indexed ticket lookups keep validation simple. This does not guarantee the required latency: SQLite serialises writes, and the 500-user target needs load testing. The local measurements are reported in the testing folder.')
    save(d,Path('2-Architecture/architecture-justification.docx'))
    d=doc('Architecture Analysis and Interface Design')
    h(d,'Comparison of the architectural styles')
    table(d,['Style','Fit for this project','Main drawback'],[
    ('Layered','Separates browser pages, business rules and storage. A single transaction can reserve a budget or consume a ticket.','A shared database can become a bottleneck. Layers still depend on well-defined interfaces.'),
    ('Microservices','Ticket validation and registration could be scaled separately if usage grew significantly.','Several deployments, network failures and cross-service consistency would add work to a small project.'),
    ('Client-server','Central records help users see the same approval state, and students only need a browser.','This describes deployment more than internal organisation. One server remains a failure point.')],[1.1,3.25,2.85])
    p(d,'Selection: a layered monolith with client-server access. These terms are compatible: layered describes the internal design, while client-server describes how the browser reaches it. There is no need for independently deployed services in this prototype.')
    h(d,'Main design decisions')
    p(d,'Budget requests advance through a fixed state sequence. A rejection ends that request. Budgets are reserved at submission, not after the Dean approves, because otherwise several pending requests could claim the same allocation. Students can register only after final approval.')
    p(d,'A ticket contains a random identifier, event ID and attendee ID, protected by an HMAC signature. Check-in compares the signature and then changes an unused ticket to used in one transaction. The database remains the source of truth for expiry and ticket use.')
    page(d);h(d,'Components and interfaces')
    table(d,['Component','Provided interface','Responsibility'],[
    ('Authentication','IAuth','Validate login, maintain sessions and enforce account roles.'),
    ('Event and Budget','IEventBudget','Create and edit drafts, reserve allocation, route approvals and keep the decision record.'),
    ('Ticket Service','ITicket','Register attendees, enforce capacity and issue signed QR tickets.'),
    ('Check-in Service','ICheckIn','Verify scanned text and consume valid tickets once.'),
    ('Dashboard Service','IStatus','Return visible events, current states, registration counts and allocation.'),
    ('SQLite Repository','IDataStore','Store users, clubs, proposals, items, decisions and tickets.'),
    ('Portal Web UI','Browser pages','Require the five business interfaces through forms and JSON requests.')],[1.65,1.25,4.3])
    p(d,'The component diagram has seven components, five UI-to-service interfaces and the shared data-access interface. Circles show provided interfaces; semicircles show required interfaces. Small boundary squares are ports. The diagram uses assembly connectors, which do not need arrowheads.')
    page(d);h(d,'Implementation mapping')
    table(d,['Logical component','Code location'],[
    ('Portal Web UI','5-Code/templates and 5-Code/static'),('Authentication','app.py: login, load_user and login_required'),('Event and Budget','services.py: create_event, submit and review'),('Ticket Service','services.py: register, token and ticket; app.py: ticket_svg'),('Check-in Service','services.py: checkin; app.py: checkin route'),('Dashboard Service','services.py: dashboard and details; app.py: status'),('SQLite Repository','services.py: db; schema.sql')],[2,5.2])
    p(d,'These are logical components, not seven separately deployed programs. The small prototype keeps related business methods in one Portal class. A larger implementation could split that class into modules without changing the basic interfaces.')
    h(d,'Data flow')
    for s in ['Proposal: the lead submits a form; the service validates it; the repository stores the draft and budget items.','Approval: the reviewer posts a decision; the service checks the current stage; the repository records the sign-off and new state together.','Ticket: a student registers; the service checks approval and capacity; the repository creates one ticket; the UI displays its signed QR.','Entry: staff submit scanned text; the service verifies it; a conditional update marks the ticket used; the page displays acceptance or rejection.']:
        p(d,s)
    h(d,'Limits')
    p(d,'The local server has no high-availability setup, university single sign-on or hardware scanner integration. Production would need HTTPS, managed credentials, backups, monitoring and a database/server setup measured under the expected load. No production deployment is claimed.')
    save(d,Path('2-Architecture/architecture-analysis.docx'))


def srs():
    d=doc('Software Requirements Specification')
    p(d,'Version 1.0. This specification describes a small portal for club event proposals, budget approvals and attendee entry. It is based on Problem Statement 10 and the revised Lab 1 requirements. It also states what the submitted prototype does not cover.')
    h(d,'1 Purpose and scope')
    p(d,'The portal keeps the event proposal, approval record and ticket registrations together. A Club Lead prepares a proposal, university reviewers decide on the budget, students register for an approved event, and staff validate tickets at entry. The system does not collect ticket payments or transfer approved money.')
    h(d,'2 Users and responsibilities')
    table(d,['User','Responsibility'],[('Club Lead','Create and edit own drafts, submit requests and track own events.'),('Faculty Coordinator','Perform the first budget review.'),('Finance Officer','Review requests cleared by the coordinator.'),('Dean','Give the final decision after finance approval.'),('Student','Register for an approved event and access their own ticket.'),('Check-in Staff','Validate QR ticket text for a selected approved event.')],[1.8,5.4])
    p(d,'Campus Admin in the original brief is represented by the reviewer and check-in roles. One account has one role. Accounts and club allocations are provisioned separately; self-service account creation is outside the prototype.')
    page(d);h(d,'3 Functional requirements')
    for rid,typ,priority,desc,criteria,rationale in REQUIREMENTS[:5]:
        p(d,rid+' ('+priority+'): '+desc)
    p(d,'The complete acceptance criteria and rationale are in requirements-and-use-case.pdf. The RTM maps each requirement to code and test cases.')
    h(d,'4 Approval states')
    p(d,'Draft -> Faculty pending -> Finance pending -> Dean pending -> Approved. At any pending stage, the assigned reviewer may choose Rejected. Approved and Rejected are final states for this prototype. A rejected request is not silently reopened; a lead can prepare a new proposal.')
    page(d);h(d,'5 Business rules')
    for s in [
    'BR01. Only the owning Club Lead may edit or submit a draft. Submitted event details are locked.',
    'BR02. Event name, date, venue, positive capacity and at least one positive budget item are required. Dates before the server\'s current date are refused.',
    'BR03. Amounts are entered in rupees with at most two decimal places and stored as integer paise. The total equals the sum of the items.',
    'BR04. Pending and approved budgets count against one configured allocation period. The check and reservation occur together. Rejection releases the reservation.',
    'BR05. Reviewers act only at their assigned stage. A sign-off records account, stage, decision and UTC time. Rejection needs a reason.',
    'BR06. A student has at most one ticket per event. Registration is limited by capacity and final approval.',
    'BR07. Tickets expire after the event date in the server\'s configured local time zone. There is no separate time-of-day entry window in this version.',
    'BR08. A valid ticket can be consumed once, even if two staff requests arrive together. Only check-in accounts may do this.',
    'BR09. Dashboard data is restricted by role and ownership. The page polls every three seconds while open.',
    'BR10. The minimum guarantee on any failed operation is no partial budget reservation, approval or ticket consumption.'
    ]: p(d,s)
    h(d,'6 Non-functional requirements')
    p(d,'NFR-001 targets complete check-in validation below 100 ms with staff authorisation, signature verification and replay rejection. The initial peak-load assumption is 20 simultaneous check-in clients.')
    p(d,'NFR-002 targets 500 concurrent users, page loads below two seconds and usable mobile forms. These are targets for acceptance; the test report records the limits of the local measurements.')
    page(d);h(d,'7 Data design')
    table(d,['Entity','Main fields and constraints'],[
    ('Club','ID, name, allocation in paise. A lead is assigned to a club.'),('User','ID, unique username, password hash, role and optional club ID.'),('Event','ID, club, owner, name, date, venue, capacity, budget and state.'),('Budget item','ID, event ID, description and positive amount.'),('Approval','ID, event ID, reviewer, stage, decision, reason and UTC time. One recorded decision per event/stage.'),('Ticket','Random ID, event ID, attendee ID and optional used time. Unique event/attendee pair.')],[1.3,5.9])
    p(d,'SQLite foreign keys enforce links between records. Parameterised SQL is used for all user values. Write operations use explicit transactions. Ticket use is updated only where used_at is still empty.')
    h(d,'8 Security and privacy')
    p(d,'The application uses password hashes, signed session cookies, CSRF tokens on state-changing forms, role checks and ticket HMAC verification. Template escaping prevents event text being interpreted as HTML. The ticket secret and database are local instance files and are excluded from Git.')
    p(d,'The demonstration accounts and password are only for local use. Internet deployment would require replacing them, enabling HTTPS and secure cookies, adding login rate limits and defining backups and retention. QR contents use internal IDs rather than names or email addresses.')
    page(d);h(d,'9 External interfaces')
    table(d,['Interface','Input and output'],[
    ('Browser pages','HTML forms for login, proposals, reviews, registration and check-in.'),('GET /api/status','Authenticated request; JSON containing visible event states, counts and remaining allocation.'),('POST /events/new or /events/{id}/edit','Lead form with event fields and budget lines; saves a draft or returns a validation message.'),('POST /events/{id}/submit','Lead request; reserves allocation and changes state.'),('POST /events/{id}/review','Reviewer decision and reason; stores sign-off or rejects an invalid stage.'),('POST /events/{id}/register','Student request; redirects to the issued or existing ticket.'),('GET /tickets/{id} and .svg','Own-ticket page and QR image; another student receives no ticket data.'),('POST /checkin','Staff account, selected event and scanned ticket text; displays acceptance or rejection.')],[2.45,4.75])
    p(d,'Every POST form includes a CSRF token. Missing or invalid tokens return HTTP 400. Unauthenticated browser requests go to login; denied pages return 403 or 404. Business-rule failures show a plain message without committing an invalid change.')
    page(d);h(d,'10 Assumptions and exclusions')
    p(d,'The university will supply the real allocation policy, account provisioning method and expected load. The prototype uses one allocation period; annual rollover and spending reconciliation are future work. An authenticated decision is the digital sign-off; certificate signing is not included.')
    p(d,'A keyboard-type QR scanner can supply ticket text to the check-in field. The current evidence uses manually pasted ticket text. There is no camera scanner, email delivery, payment gateway, attendance export or deployed university integration.')
    h(d,'11 Verification and acceptance')
    p(d,'Automated tests cover invalid proposals, approval order, rejection, budget reservation, duplicate registration, capacity, ownership, ticket tampering, replay and concurrent updates. A browser walkthrough covers the normal flow. A separate script measures local HTTP load and check-in time.')
    p(d,'Acceptance requires the functional cases to pass and the non-functional targets to be verified on the intended environment. The present prototype is a local demonstration, not a claim of production readiness. See the test report for actual results and remaining checks.')
    h(d,'12 References')
    for s in ['Problem Statement 10: 10_SE_Lab1_SE_Problem_Statements.pdf.','Git Hub Project Submission Details.docx, individual-project section.','Lab_3_Architecture_Student_handout.pdf, component diagram and justification requirements.','Revised requirements, RTM and architecture files in folders 1 and 2 of this repository.']:
        p(d,s)
    save(d,Path('4-SRS-and-Work-Breakdown/software-requirements-specification.docx'))


TASKS=[
('1.1','Read the brief and keep the original files','None','Original Lab 1 files and references','Done'),
('1.2','Revise FR and NFR acceptance criteria','1.1','Requirements document and CSV','Done'),
('1.3','Complete the RTM and use-case model','1.2','RTM, diagram and core flow','Done'),
('2.1','Compare architectures and choose one','1.2','Architecture analysis and justification','Done'),
('2.2','Draw components and define interfaces','2.1','Component diagram and code mapping','Done'),
('3.1','Prepare the SRS','1.3, 2.2','SRS document','Done'),
('3.2','Organise GitHub and capture evidence','1.1','Numbered folders and real screenshots','Done'),
('3.3','Create the Jira project and task records','3.1','Jira project and actual screenshot','Pending access'),
('4.1','Implement login, proposals and approvals','3.1','Working prototype','Done'),
('4.2','Implement tickets, check-in and status','4.1','Working prototype','Done'),
('4.3','Collect GitHub Copilot evidence','4.1','Actual Copilot session or code link','Pending tool access'),
('5.1','Test the portal, fix a defect and retest','4.2','20 tests, failure log, patch and final log','Done'),
('5.2','Measure load and review mobile layout','5.1','Benchmark results and browser evidence','Partly verified'),
('5.3','Complete the supplied game-testing exercise','Lecturer provides repository','Four supplied cases, patch and retest link','Pending input'),
('6.1','Check links and submit repository','All available tasks','Organised repository and status list','Done for available work')]


def wbs():
    d=doc('Work Breakdown and Task Plan',True)
    p(d,'This breakdown divides the individual project into small tasks. The status column describes the submitted files, not a retrospective time sheet. Jira tasks have been prepared in CSV form but have not been created in a Jira account.')
    table(d,['WBS','Task','Depends on','Output','Status'],TASKS[:8],[.6,3.15,1.15,3.1,1.7])
    page(d);h(d,'Implementation testing and submission')
    table(d,['WBS','Task','Depends on','Output','Status'],TASKS[8:],[.6,3.15,1.15,3.1,1.7])
    h(d,'Suggested sequence')
    p(d,'Finish requirements before architecture, and check the architecture before coding. Run functional tests before measuring load. Keep the original files, build the new documents, then check that README links point to the right folders. The remaining Jira, Copilot and game-exercise items depend on access or material outside this repository.')
    h(d,'Remaining work')
    p(d,'Use the provided CSV to create Jira tasks when the correct Jira project is available. Capture a real Copilot session if that tool is mandatory. When the lecturer shares the game repository, run its four supplied test cases, document the actual defect, patch it and retest. The portal regression exercise is additional practice and is not presented as that supplied game assignment.')
    save(d,Path('4-SRS-and-Work-Breakdown/work-breakdown.docx'))
    with (ROOT/'3-Project-Evidence/jira-tasks.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.writer(f);w.writerow(['Issue Type','Summary','Description','Priority','Labels'])
        for wid,task,dependency,output,status in TASKS:
            w.writerow(['Task',wid+' '+task,f'{IDENTITY}. Dependency: {dependency}. Output: {output}. Repository status: {status}.','Medium','student-club-portal'])


def testing():
    d=doc('Test Plan and Execution Report')
    p(d,'This report covers the student club portal built in this repository. It is separate from the lecturer-supplied game exercise, whose repository and four test cases were not supplied.')
    h(d,'1 Test setup')
    p(d,'Tests ran locally on Windows using Python 3.12, Flask, SQLite and pytest. Each automated test uses a fresh temporary database. The full package list and machine details are stored in results/python-packages.txt and results/performance.json. Demo data is synthetic.')
    h(d,'2 Scope and method')
    p(d,'The tests check the business services and HTTP routes. They cover both normal input and invalid actions, including concurrency for budget reservations, the last seat and duplicate entry. The browser walkthrough uses the actual running portal. No screenshot is drawn to imitate an application.')
    h(d,'3 Execution summary')
    table(d,['Run','Observed result','Evidence'],[('Initial regression run','19 passed, 1 failed','results/initial-tests.txt and initial-tests.xml'),('After the ticket fix','20 passed','results/final-tests.txt and final-tests.xml'),('Patch','Reject malformed signatures before comparison','results/ticket-validation-fix.patch')],[1.3,2.3,3.6])
    h(d,'4 Defect DEF01 Malformed ticket signature')
    p(d,'Input: a ticket string whose signature contains non-ASCII characters. Expected: the system rejects it as invalid. Actual before the fix: Python raised TypeError during hmac.compare_digest. This could produce an error page instead of the expected rejection.')
    p(d,'Fix: require a 64-character lowercase hexadecimal signature before comparing it with the expected HMAC. TC17 now passes. The full test suite was rerun to check that valid tickets and the other rejection cases still worked.')
    page(d);h(d,'5 Test cases')
    descriptions=[('TC01','Missing fields and invalid money'),('TC02','Draft and budget item storage'),('TC03','Approval order and recorded sign-offs'),('TC04','Rejection reason and allocation release'),('TC05','Over-budget request and draft revision'),('TC06','Registration before approval'),('TC07','Ticket uniqueness and repeat registration'),('TC08','Event capacity'),('TC09','Single-use entry'),('TC10','Tampering, wrong event and expiry'),('TC11','Roles and ownership'),('TC12','Concurrent scans accept once'),('TC13','Concurrent requests do not overspend'),('TC14','Status API reflects committed changes'),('TC15','Login, CSRF, logout and authentication'),('TC16','Full HTTP flow and QR response'),('TC17','Malformed non-ASCII signature'),('TC18','Concurrent registration for the last seat'),('TC19','Draft edits and locking after submission'),('TC20','HTML escaping and SQL login input')]
    table(d,['ID','Check','Final result'],[(a,b,'Passed') for a,b in descriptions],[.7,5.4,1.1])
    page(d);h(d,'6 Performance measurements')
    result=json.loads((ROOT/'6-Testing/results/performance.json').read_text())
    p(d,'Run time (UTC): '+result['run_at_utc']+'. '+result['method'])
    rows=[]
    for key,label in [('page_500_concurrent','500 concurrent page requests'),('checkin_sequential','50 sequential check-ins'),('checkin_20_workers','50 check-ins using 20 workers')]:
        r=result[key]
        rows.append((label,str(r['successful'])+'/'+str(r['requests']),str(r['median_ms']),str(r['p95_ms']),str(r['max_ms']),'Yes' if r['target_met'] else 'No'))
    table(d,['Workload','Successful','Median ms','P95 ms','Max ms','Target met'],rows,[2.2,1.0,1.0,1.0,1.0,1.0],10)
    p(d,'Targets: under 2,000 ms for each page load and under 100 ms for each check-in. These measurements include local HTTP processing but not QR hardware acquisition, browser rendering or a real network. The 500-user run uses distinct pre-authenticated users; it measures page reads, not 500 simultaneous registrations.')
    p(d,'The final run uses Waitress with 16 worker threads and direct loopback requests. Earlier runs used the development server or the default Windows proxy lookup; their raw results are retained as diagnostic baselines and are not directly comparable. The direct connection removes proxy discovery from the local measurement.')
    p(d,'All 500 final page requests succeeded, but some exceeded two seconds. Sequential check-ins met 100 ms, while the concurrent check-in run did not. NFR-001 and NFR-002 are therefore not fully met. Further work should investigate database write contention and measure the application on the intended infrastructure; the present result is a local prototype benchmark.')
    h(d,'7 Manual checks and open items')
    p(d,'Real browser screenshots record proposal creation, approval history, ticket display and check-in. A narrow viewport review checks layout only; it does not replace testing on a physical phone. The saved screenshots and steps are in 3-Project-Evidence.')
    p(d,'Open items: real scanner testing, physical mobile-device testing, browser page-load measurements on target infrastructure, and the separately assigned game repository. The GitHub Copilot and Jira evidence also remains dependent on those tools being available.')
    save(d,Path('6-Testing/test-report.docx'))


if __name__=='__main__':
    requirements();rtm();architecture();srs();wbs();testing()
